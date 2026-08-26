import io
from unittest.mock import MagicMock, patch

from django.core.files.base import ContentFile
from django.test import TestCase
from docx import Document
from pypdf import PdfWriter

from atendimento.models import AnexoAtendimento, Atendimento
from atendimento.services.attachment_extraction_service import AttachmentExtractionService


class AttachmentExtractionServiceTests(TestCase):
    def setUp(self):
        self.atendimento = Atendimento.objects.create(nome="Cliente", email="c@test.com")

    def _criar_anexo(self, nome, conteudo_bytes, tipo=AnexoAtendimento.TipoArquivo.OUTRO):
        anexo = AnexoAtendimento(
            atendimento=self.atendimento, nome_original=nome, tipo_arquivo=tipo, tamanho=len(conteudo_bytes),
        )
        anexo.arquivo.save(nome, ContentFile(conteudo_bytes), save=True)
        return anexo

    def test_docx_real_roundtrip_extracts_paragraphs_and_tables(self):
        documento = Document()
        documento.add_paragraph("Meu sistema esta fora do ar desde ontem.")
        tabela = documento.add_table(rows=1, cols=2)
        tabela.rows[0].cells[0].text = "Codigo"
        tabela.rows[0].cells[1].text = "ERR_500"
        buffer = io.BytesIO()
        documento.save(buffer)

        anexo = self._criar_anexo("relatorio.docx", buffer.getvalue())
        resultado = AttachmentExtractionService.extract(anexo)

        self.assertIsNone(resultado["erro"])
        self.assertIn("fora do ar", resultado["texto_extraido"])
        self.assertIn("ERR_500", resultado["texto_extraido"])

    def test_txt_utf8_roundtrip(self):
        anexo = self._criar_anexo("notas.txt", "Anexo em texto simples com acentuação".encode("utf-8"))
        resultado = AttachmentExtractionService.extract(anexo)
        self.assertIsNone(resultado["erro"])
        self.assertIn("acentuação", resultado["texto_extraido"])

    def test_csv_roundtrip(self):
        anexo = self._criar_anexo("dados.csv", b"servico,status\napi,fora do ar\n")
        resultado = AttachmentExtractionService.extract(anexo)
        self.assertIsNone(resultado["erro"])
        self.assertIn("fora do ar", resultado["texto_extraido"])
        self.assertIn("servico | status", resultado["texto_extraido"])

    def test_blank_pdf_is_flagged_as_possibly_scanned(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)

        anexo = self._criar_anexo("documento.pdf", buffer.getvalue())
        resultado = AttachmentExtractionService.extract(anexo)

        self.assertIsNone(resultado["erro"])
        self.assertEqual(resultado["paginas"], 1)
        self.assertFalse(resultado["disponivel"])
        self.assertIsNotNone(resultado["motivo_indisponivel"])

    @patch("atendimento.services.attachment_extraction_service.PdfReader")
    def test_pdf_with_real_text_is_marked_available(self, mock_reader_cls):
        pagina_fake = MagicMock()
        pagina_fake.extract_text.return_value = "Texto real extraido da pagina " * 5
        mock_reader_cls.return_value.pages = [pagina_fake]

        anexo = self._criar_anexo("com_texto.pdf", b"conteudo qualquer" * 10)
        resultado = AttachmentExtractionService.extract(anexo)

        self.assertIsNone(resultado["erro"])
        self.assertTrue(resultado["disponivel"])
        self.assertIn("Texto real extraido", resultado["texto_extraido"])

    def test_corrupted_pdf_never_raises(self):
        anexo = self._criar_anexo("corrompido.pdf", b"isso definitivamente nao e um pdf valido" * 5)
        resultado = AttachmentExtractionService.extract(anexo)
        self.assertIsNotNone(resultado["erro"])
        self.assertEqual(resultado["texto_extraido"], "")

    def test_unsupported_extension_returns_error_without_raising(self):
        anexo = self._criar_anexo("planilha.xlsx", b"conteudo qualquer" * 10)
        resultado = AttachmentExtractionService.extract(anexo)
        self.assertIsNotNone(resultado["erro"])

    def test_extract_all_skips_images(self):
        self._criar_anexo("notas.txt", b"texto qualquer" * 10)
        self._criar_anexo("foto.png", b"bytes de imagem qualquer" * 10, tipo=AnexoAtendimento.TipoArquivo.IMAGEM)

        resultados = AttachmentExtractionService.extract_all(self.atendimento)

        self.assertEqual(len(resultados), 1)
        self.assertEqual(resultados[0]["nome_original"], "notas.txt")
