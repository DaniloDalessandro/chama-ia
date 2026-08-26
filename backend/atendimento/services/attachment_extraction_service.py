"""
AttachmentExtractionService: extracao deterministica (zero LLM) de texto de
anexos PDF/DOCX/TXT/CSV de um atendimento. Imagens NAO passam por aqui --
vao direto para o OCRService/Agente 2.

Nunca levanta excecao: um anexo corrompido ou de extensao nao suportada so
popula o campo `erro` do resultado, preservando o anexo e nunca derrubando
a analise do atendimento.
"""

import csv
import io
import logging
import os

from docx import Document
from pypdf import PdfReader

logger = logging.getLogger(__name__)

_MIN_CHARS_POR_PAGINA_HEURISTICA = 20


def _extensao(nome: str) -> str:
    return os.path.splitext(nome)[1].lower()


def _decode_bytes(conteudo: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return conteudo.decode(encoding)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("utf-8", errors="replace")


class AttachmentExtractionService:
    @staticmethod
    def extract_all(atendimento) -> list:
        """Extrai texto de todos os anexos NAO-imagem do atendimento."""
        from ..models import AnexoAtendimento

        anexos = atendimento.anexos.exclude(tipo_arquivo=AnexoAtendimento.TipoArquivo.IMAGEM)
        return [AttachmentExtractionService.extract(anexo) for anexo in anexos]

    @staticmethod
    def extract(anexo) -> dict:
        nome = anexo.nome_original or anexo.arquivo.name
        ext = _extensao(nome)
        resultado = {
            "anexo_id": str(anexo.id),
            "nome_original": nome,
            "extensao": ext,
            "texto_extraido": "",
            "paginas": None,
            "erro": None,
            "disponivel": True,
            "motivo_indisponivel": None,
        }

        extratores = {
            ".pdf": AttachmentExtractionService._extract_pdf,
            ".docx": AttachmentExtractionService._extract_docx,
            ".txt": AttachmentExtractionService._extract_txt,
            ".csv": AttachmentExtractionService._extract_csv,
        }
        extrator = extratores.get(ext)
        if extrator is None:
            resultado["erro"] = f"Extensao nao suportada para extracao de texto: '{ext}'."
            return resultado

        try:
            extrator(anexo, resultado)
        except Exception as exc:  # noqa: BLE001 - anexo corrompido nunca deve derrubar a analise
            logger.error("Erro ao extrair anexo %s (%s): %s", anexo.id, nome, exc)
            resultado["erro"] = str(exc)
            resultado["texto_extraido"] = ""

        return resultado

    @staticmethod
    def _extract_pdf(anexo, resultado: dict) -> None:
        anexo.arquivo.open("rb")
        try:
            reader = PdfReader(anexo.arquivo)
            textos = [(page.extract_text() or "") for page in reader.pages]
        finally:
            anexo.arquivo.close()

        n_paginas = len(textos)
        resultado["paginas"] = n_paginas
        texto_total = "\n".join(textos)
        resultado["texto_extraido"] = texto_total

        media_chars = (len(texto_total) / n_paginas) if n_paginas else 0
        if media_chars < _MIN_CHARS_POR_PAGINA_HEURISTICA:
            resultado["disponivel"] = False
            resultado["motivo_indisponivel"] = (
                "PDF parece ser escaneado (pouco texto extraivel por pagina); renderizacao "
                "de pagina para OCR nao esta disponivel neste ambiente."
            )

    @staticmethod
    def _extract_docx(anexo, resultado: dict) -> None:
        anexo.arquivo.open("rb")
        try:
            documento = Document(anexo.arquivo)
        finally:
            anexo.arquivo.close()

        partes = [paragrafo.text for paragrafo in documento.paragraphs if paragrafo.text.strip()]
        for tabela in documento.tables:
            for linha in tabela.rows:
                celulas = [celula.text.strip() for celula in linha.cells]
                partes.append(" | ".join(celulas))

        resultado["texto_extraido"] = "\n".join(partes)

    @staticmethod
    def _extract_txt(anexo, resultado: dict) -> None:
        anexo.arquivo.open("rb")
        try:
            conteudo = anexo.arquivo.read()
        finally:
            anexo.arquivo.close()
        resultado["texto_extraido"] = _decode_bytes(conteudo)

    @staticmethod
    def _extract_csv(anexo, resultado: dict) -> None:
        anexo.arquivo.open("rb")
        try:
            conteudo = anexo.arquivo.read()
        finally:
            anexo.arquivo.close()

        texto = _decode_bytes(conteudo)
        linhas = [" | ".join(linha) for linha in csv.reader(io.StringIO(texto))]
        resultado["texto_extraido"] = "\n".join(linhas)
